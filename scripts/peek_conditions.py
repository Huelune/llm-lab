"""Polyspace 결과 엑셀의 조건 시트(MISRA 규칙, Code Metrics, RTE)를 한 화면에 보여 준다.

엑셀은 읽기만 한다.

    uv run python scripts/peek_conditions.py "엑셀 경로.xlsx" [misra] [metric] [rte]

이름에 misra·metric·rte가 들어가고 _Result로 끝나지 않는 시트를 조건 시트로 본다. 뒤에 낱말을 주면
그 낱말이 이름에 든 조건 시트만 보여 준다. 병합 셀이 남긴 빈 열은 빼고, 빈칸은 `-`로 찍는다.
행이 화면에 다 들어가지 않는 시트만 값 종류 줄을 함께 찍는다.

화면은 캡처해서 회사 밖(Claude)으로 보낸다. 그래서 칸 값은 숫자이거나 공개된 표준 용어
(Polyspace 지표·검사 이름, MISRA 분류, 색 등)로만 이뤄졌을 때만 그대로 찍고, 나머지 글은
`<글 N자>`처럼 길이만 찍는다. 엑셀 파일 이름, 조건 시트가 아닌 시트 이름, 오류 메시지 원문도
찍지 않는다. 임시 파일도 쓰지 않는다.
"""

from __future__ import annotations

import os
import sys
import unicodedata
import warnings
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from llm_lab.conditions import Metric, read_metrics
from llm_lab.safe_text import is_number, safe

KEYWORDS = ("misra", "metric", "rte")
HEADER_SCAN_ROWS = 30
WIDTH = 118
SCREEN_LINES = 40
FEW_VALUES = 8  # 값 종류가 이 이하이고 겹치는 값이 있는 열만 종류별 개수를 보여 준다
MAX_KINDS = 3  # 시트마다 화면에 찍는 값 종류 줄 (조건 시트가 4개 이하일 때)

# 오류는 종류만 알린다. 원문에는 엑셀 경로나 칸 값이 들어 있을 수 있다.
ERROR_HINTS = {
    "FileNotFoundError": "엑셀 파일이 없습니다. 경로를 확인하세요.",
    "PermissionError": "엑셀을 읽을 수 없습니다. 엑셀 프로그램에서 열려 있으면 닫고 다시 하세요.",
    "InvalidFileException": ".xlsx 파일이 아닙니다.",
    "BadZipFile": ".xlsx 파일이 아니거나 깨졌습니다.",
}

Row = tuple[int, list[str]]  # (엑셀 행 번호, 칸 글자), 빈 행은 없다


def text(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def width(line: str) -> int:
    """콘솔에서 차지하는 칸 수 (한글은 2칸)."""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line)


def cut(line: str) -> str:
    """WIDTH 칸에 맞게 자른다('..'). 한글 콘솔에서 2칸으로 그려지는 '…'는 쓰지 않는다."""
    if width(line) <= WIDTH:
        return line
    while width(line) > WIDTH - 2:
        line = line[:-1]
    return line + ".."


def wordy(value: str) -> bool:
    """숫자가 아닌 글자 칸인지 (헤더 추정용)."""
    return bool(value) and not is_number(value)


def is_condition(title: str, keys: tuple[str, ...] = KEYWORDS) -> bool:
    name = "".join(ch for ch in title.casefold() if ch.isalnum())
    return not title.casefold().endswith("_result") and any(key in name for key in keys)


def read(sheet: Any) -> list[Row]:
    rows: list[Row] = []
    for number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [text(cell) for cell in cells]
        while values and not values[-1]:
            values.pop()
        if values:
            rows.append((number, values))
    return rows


def compact(rows: list[Row]) -> tuple[list[Row], int]:
    """모든 행에서 비어 있는 열(병합 셀이 남긴 빈칸)을 뺀다. (행, 뺀 열 수)."""
    total = max((len(values) for _, values in rows), default=0)
    used = [i for i in range(total) if any(i < len(v) and v[i] for _, v in rows)]
    packed: list[Row] = []
    for number, values in rows:
        kept = [values[i] if i < len(values) else "" for i in used]
        while kept and not kept[-1]:
            kept.pop()
        packed.append((number, kept))
    return packed, total - len(used)


def kinds(rows: list[Row]) -> tuple[int, list[str]]:
    """(헤더로 보이는 행의 엑셀 번호, 값 종류 줄).

    처음 행들 중 글자 칸이 가장 많은 행을 헤더로 본다.
    """
    if not rows:
        return 0, []
    scan = rows[:HEADER_SCAN_ROWS]
    top = max(range(len(scan)), key=lambda i: (sum(map(wordy, scan[i][1])), -i))
    names = scan[top][1]
    data = [values for _, values in rows[top + 1 :]]
    found: list[tuple[int, int, str]] = []
    for i, name in enumerate(names):
        counter = Counter(values[i] if i < len(values) else "" for values in data)
        if 1 < len(counter) <= FEW_VALUES and len(counter) < len(data):
            listed = ", ".join(
                f"{safe(value) or '(빈칸)'} {n}" for value, n in counter.most_common()
            )
            found.append((len(counter), i, f"{safe(name) or f'(열{i + 1})'}: {listed}"))
    return scan[top][0], [line for _, _, line in sorted(found)]


def row_line(row: Row) -> str:
    number, values = row
    return f"  {number:>4}: " + " | ".join(safe(value) or "-" for value in values)


def share(needs: list[int], budget: int) -> list[int]:
    """줄 수를 나눈다. 적게 필요한 시트부터 다 주고 남는 줄을 나머지가 고르게 쓴다."""
    given = [0] * len(needs)
    order = sorted(range(len(needs)), key=lambda i: needs[i])
    for k, i in enumerate(order):
        given[i] = min(needs[i], max(budget, 0) // (len(order) - k))
        budget -= given[i]
    return given


def metric_line(metrics: list[Metric], note: str) -> str:
    """llm-fix-server가 Code Metrics 조건을 어떻게 읽는지 개수로만 (이름·글은 찍지 않는다)."""
    if not metrics:
        return f"Code Metrics 조건 읽기: {note}"
    bands = [band for metric in metrics for band in metric.bands]
    kinds_ = Counter(band.kind for band in bands)
    unread = sum(not band.ok for band in bands)
    return (
        f"Code Metrics 조건 읽기: 지표 {len(metrics)}개, 단계 {len(bands)}개 "
        f"(pass {kinds_['pass']} / middle {kinds_['middle']} / fail {kinds_['fail']}), "
        f"범위 글 못 읽음 {unread}"
    )


def summarize(path: Path, only: tuple[str, ...] = ()) -> list[str]:
    """화면 줄 (SCREEN_LINES - 1 이하).

    엑셀을 열지 못하면 예외를 그대로 낸다 (main이 종류만 알린다).
    """
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = [(sheet.title, read(sheet)) for sheet in workbook.worksheets]
            reading = metric_line(*read_metrics(workbook))
        finally:
            workbook.close()

    conditions = [(title, rows) for title, rows in sheets if is_condition(title)]
    others = len(sheets) - len(conditions)
    found = [(title, rows) for title, rows in conditions if not only or is_condition(title, only)]
    if not found:
        lines = [
            "조건 시트를 찾지 못했습니다. 이름에 "
            + ", ".join(only or KEYWORDS)
            + " 중 하나가 들어가고 _Result로 끝나지 않는 시트를 찾습니다.",
            f"시트 {len(sheets)}개 (표준 용어가 아닌 이름은 길이만):",
            *(f"  [{safe(title)}] {len(rows)}행" for title, rows in sheets),
        ]
        return [cut(line) for line in lines[: SCREEN_LINES - 1]]

    picked = (
        f"{len(conditions)}개 중 {len(found)}개 ({', '.join(only)})" if only else f"{len(found)}개"
    )
    head = [f"조건 시트 {picked} / 그 밖의 시트 {others}개 (이름은 찍지 않음)", reading]
    limit = MAX_KINDS if len(found) <= 4 else 0
    described = [(title, *compact(rows)) for title, rows in found]
    guessed = [kinds(rows) for _, rows, _ in described]
    needs = [len(rows) for _, rows, _ in described]
    budget = SCREEN_LINES - 1 - len(head) - len(described)  # 시트마다 제목 한 줄
    # 행이 다 들어가지 않는 시트만 값 종류 줄을 찍는다
    cut_off = [given < need for given, need in zip(share(needs, budget), needs, strict=True)]
    shown_kinds = [k[:limit] if c else [] for (_, k), c in zip(guessed, cut_off, strict=True)]
    given = share(needs, budget - sum(map(len, shown_kinds)))

    screen = list(head)
    for (title, rows, dropped), (top, _), extra, shown in zip(
        described, guessed, shown_kinds, given, strict=True
    ):
        columns = max((len(values) for _, values in rows), default=0)
        guess = f"헤더 추정 {top}행 / 열 {columns}개" if top else "비어 있음"
        if dropped:
            guess += f" (빈 열 {dropped}개 뺌)"
        screen.append(f"[{safe(title)}] {len(rows)}행 중 앞 {shown}행 / {guess}")
        screen += [f"  종류 {line}" for line in extra]
        screen += [row_line(row) for row in rows[:shown]]
    return [cut(line) for line in screen]


def clear_screen() -> None:
    # 고정 명령(cls/clear)만 쓰므로 os.system으로 충분하다
    os.system("cls" if os.name == "nt" else "clear")


def main() -> int:
    only = tuple(word.casefold() for word in sys.argv[2:])
    if len(sys.argv) < 2 or any(word not in KEYWORDS for word in only):
        print('사용법: uv run python scripts/peek_conditions.py "엑셀.xlsx" [misra|metric|rte]...')
        return 2
    path = Path(sys.argv[1].strip().strip('"'))
    try:
        screen = summarize(path, only)
    except Exception as exc:
        kind = type(exc).__name__
        clear_screen()
        print(f"엑셀을 읽지 못했습니다 (오류 종류: {kind}).")
        print(ERROR_HINTS.get(kind, "엑셀 프로그램에서 열리는 .xlsx 파일인지 확인하세요."))
        return 1
    clear_screen()
    print("\n".join(screen))
    return 0


if __name__ == "__main__":
    sys.exit(main())
