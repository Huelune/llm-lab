"""Polyspace 결과 엑셀의 시트 모양을 한 화면에 보여 준다 (MISRA 준비용, 읽기만 한다).

    uv run python scripts/peek_excel.py "엑셀 경로.xlsx"

화면에는 시트별 헤더·값 종류·예시 1행만 찍고, 예시 5행은 임시 폴더 파일에 쓴다.
"""

from __future__ import annotations

import os
import sys
import tempfile
import warnings
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HEADER_SCAN_ROWS = 30
WIDTH = 118
FEW_VALUES = 8  # 값 종류가 이 이하인 열만 종류별 개수를 보여 준다
FILE_EXAMPLES = 5


def text(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def cut(line: str, width: int = WIDTH) -> str:
    return line if len(line) <= width else line[: width - 1] + "…"


def header_row(rows: list[tuple[Any, ...]]) -> int:
    """처음 행들 중 글자 칸이 가장 많은 행을 헤더로 본다 (0부터)."""
    counts = [sum(isinstance(cell, str) and cell.strip() != "" for cell in row) for row in rows]
    return max(range(len(counts)), key=lambda i: (counts[i], -i)) if counts else 0


def describe(sheet: Any) -> tuple[list[str], list[str]]:
    """(화면 줄, 파일 줄)."""
    rows = [row for row in sheet.iter_rows(values_only=True)]
    if not rows:
        return [f"[{sheet.title}] 비어 있음"], []
    top = header_row(rows[:HEADER_SCAN_ROWS])
    names = [text(cell) or f"(열{i + 1})" for i, cell in enumerate(rows[top])]
    data = [row for row in rows[top + 1 :] if any(text(cell) for cell in row)]
    screen = [
        f"[{sheet.title}] 헤더 {top + 1}행, 데이터 {len(data)}행, 열 {len(names)}개",
        cut("  열: " + " | ".join(names)),
    ]
    few, lengths = [], []
    for i, name in enumerate(names):
        values = [text(row[i]) if i < len(row) else "" for row in data]
        counter = Counter(values)
        longest = max((len(value) for value in values), default=0)
        lengths.append(f"{name}={longest}")
        if 1 < len(counter) <= FEW_VALUES:
            kinds = ", ".join(f"{value or '(빈칸)'} {n}" for value, n in counter.most_common())
            few.append(f"{name}: {kinds}")
    screen += [cut("  종류 " + line) for line in few[:5]]
    screen.append(cut("  최대 글자 수: " + ", ".join(lengths)))
    if data:
        screen.append("  예시 1행:")
        screen += [
            cut(f"    {name}: {text(cell)}", WIDTH)
            for name, cell in zip(names, data[0], strict=False)
            if text(cell)
        ]
    full = [f"[{sheet.title}] 헤더 {top + 1}행", "열: " + " | ".join(names)] + few
    for number, row in enumerate(data[:FILE_EXAMPLES], start=1):
        full.append(f"예시 {number}:")
        full += [f"  {n}: {text(c)}" for n, c in zip(names, row, strict=False) if text(c)]
    return screen, full + [""]


def main() -> int:
    if len(sys.argv) != 2:
        print('사용법: uv run python scripts/peek_excel.py "엑셀 경로.xlsx"')
        return 2
    path = Path(sys.argv[1].strip('"'))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            print(f"엑셀을 열 수 없습니다: {path}\n이유: {exc}")
            return 1
        try:
            targets = [s for s in workbook.worksheets if s.title.casefold().endswith("_result")]
            others = [s.title for s in workbook.worksheets if s not in targets]
            screen: list[str] = []
            full: list[str] = []
            for sheet in targets:
                if "rte" in sheet.title.casefold():
                    screen.append(f"[{sheet.title}] RTE 시트라 건너뜀")
                    continue
                shown, written = describe(sheet)
                screen += shown
                full += written
        finally:
            workbook.close()
    out = Path(tempfile.gettempdir()) / "peek_excel.txt"
    out.write_text("\n".join(full), encoding="utf-8")
    os.system("cls" if os.name == "nt" else "clear")
    print(cut(f"파일: {path.name}  _Result 시트 {len(targets)}개"))
    print(cut("그 밖의 시트: " + (", ".join(others) or "없음")))
    for line in screen:
        print(line)
    print(f"예시 {FILE_EXAMPLES}행씩 전체: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
