"""시스템 구조도 한 장을 편집 가능한 PowerPoint 도형으로 만든다."""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("시스템_구조도.pptx")
FONT = "맑은 고딕"
S = 0.0143  # 도면 단위 -> 인치


def X(v):
    return Inches(0.5 + v * S)


def Y(v):
    return Inches(0.42 + v * S)


def L(v):
    return Inches(v * S)


INK, QUIET, EDGE = (
    RGBColor(0x1F, 0x1F, 0x1F),
    RGBColor(0x5F, 0x5F, 0x5F),
    RGBColor(0xA8, 0xA8, 0xA0),
)
ACCENT, ACCENT_FILL = RGBColor(0x24, 0x63, 0xD1), RGBColor(0xE4, 0xEC, 0xFA)
TINT, WHITE = RGBColor(0xF0, 0xF0, 0xEC), RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 레이아웃
shapes = slide.shapes


def set_font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def write(shape, lines, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.03):
    """lines = [(글, 크기, 색, 굵게), ...] 한 줄에 한 문단."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        set_font(p.add_run(), size, color, bold)
        p.runs[0].text = text


def rect(x, y, w, h, fill, line, width=1.0, dash=False, radius=0.06, name=""):
    shp = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x), Y(y), L(w), L(h))
    shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(width)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    if name:
        shp.name = name
    return shp


def text(x, y, w, h, value, size, color, bold=True, align=PP_ALIGN.CENTER):
    tb = shapes.add_textbox(X(x), Y(y), L(w), L(h))
    write(tb, [(value, size, color, bold)], align=align, margin=0)
    return tb


def box(x, y, w, h, name, lines, main=False, name_size=14):
    shp = rect(
        x,
        y,
        w,
        h,
        ACCENT_FILL if main else WHITE,
        ACCENT if main else EDGE,
        2 if main else 1,
        name=name,
    )
    write(
        shp,
        [(name, name_size, INK, True)] + [(t, 11, INK if main else QUIET, False) for t in lines],
    )


def db(x, name, sub):
    shp = shapes.add_shape(MSO_SHAPE.CAN, X(x), Y(294), L(88), L(72))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = EDGE
    shp.line.width = Pt(1)
    shp.name = name
    write(shp, [(name, 12, INK, True), (sub, 10.5, QUIET, False)], anchor=MSO_ANCHOR.BOTTOM)
    shp.text_frame.margin_bottom = Inches(0.08)


def arrow(x1, y1, x2, y2, both=False):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(x1), Y(y1), X(x2), Y(y2))
    c.line.color.rgb = EDGE
    c.line.width = Pt(1.25)
    ln = c.line._get_or_add_ln()
    if both:
        etree.SubElement(ln, qn("a:headEnd")).set("type", "triangle")
    etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")


# 제목
tb = shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.5))
write(
    tb,
    [("엑셀과 소스를 넣으면 회사 PC 안에서 패치가 나온다", 24, INK, True)],
    align=PP_ALIGN.LEFT,
    margin=0,
)

# 열 배경과 Web 영역
rect(16, 44, 180, 356, TINT, None, radius=0.05, name="입력 열")
rect(664, 44, 180, 356, TINT, None, radius=0.05, name="연결·결과 열")
rect(212, 44, 436, 356, WHITE, ACCENT, 1.25, dash=True, radius=0.04, name="Web 영역")
text(16, 52, 180, 22, "입력", 13, QUIET)
text(212, 52, 436, 22, "Web · 회사 PC 한 대 안에서 돔", 13, QUIET)
text(664, 52, 180, 22, "연결 · 결과", 13, QUIET)

# 입력
box(
    28,
    168,
    156,
    92,
    "Polyspace 결과",
    ["RTE · MISRA C:2012", "Code Metrics", "+ 조건 시트(회사 기준)"],
)
box(28, 294, 156, 72, "제어 SW 소스", ["프로젝트 전체 .c / .h"])

# Web
box(228, 84, 120, 64, "실행기", ["명령 하나로 실행", "이 PC에서만 접속"])
box(364, 84, 268, 64, "작업 엔진", ["중요한 것부터 동시 처리", "진행률 · 멈춤 · 이어서"], main=True)
box(228, 172, 88, 84, "① 결과 수집", ["시트·조건 읽기", "소스 짝짓기"], name_size=12.5)
box(333, 172, 88, 84, "② 문맥 찾기", ["정의·호출자", "묶음 만들기"], name_size=12.5)
box(438, 172, 88, 84, "③ LLM 판단", ["수정 여부·이유", "되묻기"], name_size=12.5)
box(543, 172, 88, 84, "④ 검사·패치", ["위치·겹침", "패치 생성"], name_size=12.5)
db(228, "소스 원본", "패치 기준")
db(333, "정의·호출자", "프로젝트 색인")
db(438, "LLM 답·기록", "재사용·원문")
db(543, "작업·결과", "행별 판단·diff")
text(228, 374, 403, 18, "로컬 DB (SQLite) · 회사 PC 안에만 저장", 12, ACCENT)

# 연결·결과
box(676, 84, 156, 64, "회사 LLM", ["사내망 전용", "유일한 바깥 연결"], main=True)
box(676, 172, 156, 84, "웹 화면", ["수정 / 불필요 / 직접 수정", "diff 보고 고르기"])
box(676, 294, 156, 72, "패치 파일", [".c / .h 개선안", "+ 코멘트용 이유"])

# 화살표
arrow(348, 116, 364, 116)
arrow(632, 108, 676, 108)
arrow(676, 124, 632, 124)
for x in (288, 377, 482, 587):
    arrow(x, 148, x, 172)
for x1, x2 in ((316, 333), (421, 438), (526, 543)):
    arrow(x1, 214, x2, 214)
arrow(184, 214, 228, 214)
arrow(184, 324, 228, 240)
arrow(631, 214, 676, 214)
for x in (272, 377, 482, 587):
    arrow(x, 256, x, 294)
arrow(631, 330, 676, 330)

# 데이터 흐름
text(16, 432, 100, 28, "데이터 흐름", 15, INK, align=PP_ALIGN.LEFT)
flow = [
    ("엑셀·소스 입력", "Polyspace 결과"),
    ("문맥 수집", "정의·호출자 묶기"),
    ("LLM 판단", "필요하면 되묻기"),
    ("검토·선택", "사람이 diff 확인"),
    ("패치·재분석", "Polyspace로 확인"),
]
for i, (name, sub) in enumerate(flow):
    x = 120 + i * 146
    shp = rect(x, 420, 128, 52, ACCENT_FILL, ACCENT, 1.25, name=f"흐름 {i + 1}")
    write(shp, [(name, 13, INK, True), (sub, 11, QUIET, False)])
    if i < 4:
        arrow(x + 128, 446, x + 146, 446)

slide.notes_slide.notes_text_frame.text = "\n".join(
    [
        "넣는 것: Polyspace 결과 엑셀(RTE, MISRA C:2012, Code Metrics, 조건 시트)과 "
        "제어 SW 소스 폴더.",
        "Web에서 하는 일: 결과 수집 → 문맥 찾기 → LLM 판단 → 검사·패치. "
        "작업 엔진이 중요한 것부터 동시에 돌린다.",
        "쓰는 DB: SQLite 하나. 소스 원본, 정의·호출자 색인, LLM 답과 대화 기록, "
        "작업·행별 결과를 회사 PC 안에만 저장.",
        "결과: 웹 화면(수정 / 수정 불필요 / 직접 수정)과 "
        "패치 파일(.c / .h 개선안 + 코멘트용 이유).",
        "바깥 연결은 사내망의 회사 LLM 하나뿐이다.",
    ]
)

prs.save(OUT)
print("written")
