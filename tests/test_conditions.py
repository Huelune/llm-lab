"""conditions: Polyspace 엑셀의 Code Metrics 조건 시트에서 지표별 단계를 읽는다.

조건 값은 바뀔 수 있으므로 실제 값이 아니라 지어낸 지표 이름과 범위를 쓴다.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from llm_lab.conditions import Band, Metric, parse_range, read_metrics, verdict_kind


def _row(cells: dict[int, Any], width: int = 15) -> list[Any]:
    row: list[Any] = [None] * width
    for index, value in cells.items():
        row[index] = value
    return row


def _merged_sheet(book: Workbook, title: str = "Fake Metrics") -> None:
    """병합 셀이 남긴 빈 열, 위쪽 빈 행, 첫 단계 행에만 있는 이름."""
    sheet = book.create_sheet(title)
    sheet.append([])
    sheet.append(["지표 기준표"])
    sheet.append(_row({1: "No", 2: "Metric Name", 5: "설명", 12: "기준 :", 14: "PASS/FAIL"}))
    sheet.append(_row({1: 1, 2: "Fake Depth", 5: "가짜 설명", 12: "2 ~ 9", 14: "Pass"}))
    sheet.append(_row({12: "10~19", 14: "검토 필요"}))
    sheet.append(_row({12: ">=20", 14: "Fail (수정 필요)"}))
    sheet.append(_row({1: 2, 2: "Fake Ratio", 12: "≥ 40%", 14: "pass"}))
    sheet.append(_row({12: "< 40", 14: "FAIL"}))
    sheet.append([])
    sheet.append(_row({1: 3, 2: "Fake Count", 12: "7 이하", 14: "Pass"}))
    sheet.append(_row({12: "많음", 14: "Fail"}))


def _save(book: Workbook, path: Path) -> Any:
    book.save(path)
    return load_workbook(path, read_only=True, data_only=True)


def _book() -> Workbook:
    book = Workbook()
    book.active.title = "Cover"
    book.active.append(["표지"])
    misra = book.create_sheet("Fake Rules")  # MISRA 조건 시트처럼 범위·판정 열이 없다
    misra.append(["Guideline", "Mode", "Enabled"])
    misra.append(["1.1", "required", "yes"])
    result = book.create_sheet("CodeMetrics_Result")  # 결과 시트는 보지 않는다
    result.append(["File", "Function", "Check", "Threshold", "Pass / Fail"])
    return book


def test_reads_bands_through_merged_cells(tmp_path: Path) -> None:
    book = _book()
    _merged_sheet(book)

    metrics, note = read_metrics(_save(book, tmp_path / "c.xlsx"))

    assert [m.name for m in metrics] == ["Fake Depth", "Fake Ratio", "Fake Count"]
    depth = metrics[0]
    assert [(b.low, b.high, b.kind) for b in depth.bands] == [
        (2, 9, "pass"),
        (10, 19, "middle"),
        (20, None, "fail"),
    ]
    assert [b.verdict for b in depth.bands] == ["Pass", "검토 필요", "Fail (수정 필요)"]
    assert depth.pass_band is not None and depth.pass_band.text == "2 ~ 9"
    assert note == "조건 시트: 지표 3개 (범위 글을 읽지 못한 단계 1개)"


def test_band_contains_follows_open_and_closed_ends(tmp_path: Path) -> None:
    book = _book()
    _merged_sheet(book)
    depth, ratio, count = read_metrics(_save(book, tmp_path / "c.xlsx"))[0]

    assert [depth.stage(v).kind for v in (2, 9, 10, 19, 20, 500)] == [  # type: ignore[union-attr]
        "pass",
        "pass",
        "middle",
        "middle",
        "fail",
        "fail",
    ]
    assert depth.stage(1) is None and depth.stage(9.5) is None  # 범위 밖
    # 낮을수록 나쁜 지표도 방향 없이 범위로만 판정한다
    assert ratio.stage(40).kind == "pass" and ratio.stage(39.9).kind == "fail"  # type: ignore[union-attr]
    # 읽지 못한 범위 글("많음")은 어떤 값도 담지 않는다
    assert count.stage(7).kind == "pass" and count.stage(8) is None  # type: ignore[union-attr]
    assert not count.bands[1].ok


def test_picks_sheet_by_columns_and_prefers_metric_in_title(tmp_path: Path) -> None:
    book = _book()
    other = book.create_sheet("Other Table")  # 열은 맞지만 이름에 metric이 없다
    other.append(["No", "Metric", "기준", "Pass / Fail"])
    other.append([1, "Fake Other", "0 ~ 1", "Pass"])
    _merged_sheet(book, "Code-Metrics (2)")

    metrics, _ = read_metrics(_save(book, tmp_path / "c.xlsx"))

    assert [m.name for m in metrics] == ["Fake Depth", "Fake Ratio", "Fake Count"]


def test_any_sheet_with_the_columns_is_used(tmp_path: Path) -> None:
    book = _book()
    other = book.create_sheet("Other Table")
    other.append(["No", "Metric", "기준", "Pass / Fail"])
    other.append([1, "Fake Other", "0 ~ 1", "Pass"])

    metrics, _ = read_metrics(_save(book, tmp_path / "c.xlsx"))

    assert [m.name for m in metrics] == ["Fake Other"]


def test_name_column_falls_back_to_right_of_no(tmp_path: Path) -> None:
    book = _book()
    sheet = book.create_sheet("Fake")
    sheet.append(["No", "가짜머리글", "기준", "판정"])
    sheet.append([1, "Fake Paths", "0 ~ 3", "Pass"])
    sheet.append([None, None, "> 3", "Fail"])

    metrics, _ = read_metrics(_save(book, tmp_path / "c.xlsx"))

    assert [m.name for m in metrics] == ["Fake Paths"]
    assert metrics[0].bands[1].low == 3 and metrics[0].bands[1].low_open


def test_missing_sheet_gives_empty_result_and_reason(tmp_path: Path) -> None:
    metrics, note = read_metrics(_save(_book(), tmp_path / "c.xlsx"))

    assert metrics == []
    assert note == "조건 시트 없음: 결과 시트 Threshold만 씀"


def test_sheet_without_metric_rows_gives_reason(tmp_path: Path) -> None:
    book = _book()
    book.create_sheet("Metrics").append(["No", "Metric", "기준", "Pass / Fail"])

    metrics, note = read_metrics(_save(book, tmp_path / "c.xlsx"))

    assert metrics == []
    assert note == "조건 시트에 읽을 지표 없음: 결과 시트 Threshold만 씀"


def test_round_trips_through_json(tmp_path: Path) -> None:
    book = _book()
    _merged_sheet(book)
    metrics, _ = read_metrics(_save(book, tmp_path / "c.xlsx"))

    data = [m.to_json() for m in metrics]
    again = [Metric.from_json(item) for item in data]

    assert again == metrics
    assert again[0].stage(15).kind == "middle"  # type: ignore[union-attr]


@pytest.mark.parametrize(
    ("text", "parsed"),
    [
        ("1 ~ 15", (1, 15, False, False)),
        ("16~30", (16, 30, False, False)),
        ("3-8", (3, 8, False, False)),
        (">= 31", (31, None, False, False)),
        (">=11", (11, None, False, False)),
        ("=> 4", (4, None, False, False)),
        ("> 5", (5, None, True, False)),
        ("<= 2.5", (None, 2.5, False, False)),
        ("< 10%", (None, 10, False, True)),
        ("≤ 7", (None, 7, False, False)),
        ("20 이상", (20, None, False, False)),
        ("20 미만", (None, 20, False, True)),
        ("1,000 초과", (1000, None, True, False)),
        ("0", (0, 0, False, False)),
        ("약 10", None),
        ("", None),
    ],
)
def test_parse_range(text: str, parsed: tuple[Any, ...] | None) -> None:
    assert parse_range(text) == parsed


@pytest.mark.parametrize(
    ("verdict", "kind"),
    [
        ("Pass", "pass"),
        ("Fail (수정 필요)", "fail"),
        ("통과", "pass"),
        ("불합격", "fail"),
        ("검토 필요", "middle"),
        ("Pass / Fail", "middle"),
    ],
)
def test_verdict_kind(verdict: str, kind: str) -> None:
    assert verdict_kind(verdict) == kind


def test_band_is_frozen() -> None:
    band = Band(1, 2, False, False, "1 ~ 2", "Pass", "pass")
    with pytest.raises(AttributeError):
        band.low = 3  # type: ignore[misc]
