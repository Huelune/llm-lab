"""scripts/peek_metrics.py: CodeMetrics 시트 요약이 한 화면에 맞고 숫자가 맞는지."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path
from types import ModuleType

from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "peek_metrics.py"
HEADER = ["File", "Function", "line", "Check", "Threshold", "Actual Value", "Status", "Comment"]


def _load() -> ModuleType:
    spec = importlib.util.spec_from_file_location("peek_metrics", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # dataclass가 모듈을 찾는다
    spec.loader.exec_module(module)
    return module


peek = _load()


def _workbook(path: Path, rows: list[list[object]], title: str = "CodeMetrics_Result") -> Path:
    book = Workbook()
    book.active.title = "Summary"
    rte = book.create_sheet("RTE_Result")
    rte.append(["ID", "TYPE", "File", "line", "check", "detail"])
    sheet = book.create_sheet(title)
    sheet.append(["Polyspace Code Metrics"])  # 제목 행
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _check_line(screen: list[str], name: str) -> tuple[list[str], list[str]]:
    """(앞 두 숫자: 행·함수 빈칸, 뒤 네 숫자: >기준 =기준 <기준 비교불가)."""
    line = next(line for line in screen if line.startswith(name + " "))
    tokens = line[len(name) :].split()
    return tokens[:2], tokens[-4:]


ROWS: list[list[object]] = [
    ["src/a.c", "foo", 10, "Cyclomatic Complexity", 10, 15, "Unreviewed", None],
    ["src/a.c", "foo", 10, "Number of Paths", 80, 120, "Unreviewed", None],
    ["src/a.c", "bar", 50, "Cyclomatic Complexity", "10", "12", "Justified", "ok"],
    [None, None, None, None, None, None, None, None],  # 빈 행은 세지 않는다
    ["src/b.c", None, None, "Comment Density", 20, 5, "Unreviewed", None],
    ["src/b.c", "baz", 7, "Number of Paths", "<= 80", 90, "Unreviewed", None],
]


def test_counts_per_check(tmp_path: Path) -> None:
    screen, full = peek.summarize(_workbook(tmp_path / "m.xlsx", ROWS))

    assert any("CodeMetrics_Result" in line and "데이터 5행" in line for line in screen)
    assert _check_line(screen, "Cyclomatic Complexity") == (["2", "0"], ["2", "0", "0", "0"])
    # 숫자가 아닌 Threshold("<= 80")는 비교하지 않고 따로 센다
    assert _check_line(screen, "Number of Paths") == (["2", "0"], ["1", "0", "0", "1"])
    assert _check_line(screen, "Comment Density") == (["1", "1"], ["0", "0", "1", "0"])
    paths = next(line for line in screen if line.startswith("Number of Paths"))
    assert "80" in paths and "<= 80" in paths and "90~120" in paths


def test_function_and_status_lines(tmp_path: Path) -> None:
    screen, _ = peek.summarize(_workbook(tmp_path / "m.xlsx", ROWS))
    text = "\n".join(screen)

    assert "함수(File+Function) 3개" in text
    assert "함수당 행 수: 1개 2 / 2개 1" in text
    assert "Unreviewed 4" in text and "Justified 1" in text
    assert "Comment 있는 행 1" in text
    assert "Function 빈 행 1" in text


def test_line_differs_within_function(tmp_path: Path) -> None:
    rows: list[list[object]] = [
        ["a.c", "foo", 10, "Cyclomatic Complexity", 10, 15, "Unreviewed", None],
        ["a.c", "foo", 12, "Number of Paths", 80, 120, "Unreviewed", None],
    ]
    screen, _ = peek.summarize(_workbook(tmp_path / "m.xlsx", rows))

    assert "line이 다른 함수 1개" in "\n".join(screen)
    # 보통 크기면 아무 줄도 잘리지 않는다
    assert not any(line.endswith("..") for line in screen)


def test_many_checks_fit_one_screen(tmp_path: Path) -> None:
    rows: list[list[object]] = [
        [
            "a.c",
            f"f{i}",
            i,
            f"측정 지표 이름이 아주 긴 한글 지표 {i:02d} 번째",
            10,
            11,
            "Unreviewed",
            None,
        ]
        for i in range(60)
    ]
    screen, full = peek.summarize(_workbook(tmp_path / "m.xlsx", rows))

    assert len(screen) <= peek.SCREEN_LINES
    assert all(peek.width(line) <= peek.WIDTH for line in screen)
    assert any("외 " in line and "종" in line for line in screen)
    # 파일에는 잘린 지표까지 모두 들어간다
    assert sum("지표 59 번째" in line for line in full) >= 1


def test_missing_sheet_explains_why(tmp_path: Path) -> None:
    book = Workbook()
    book.active.title = "RTE_Result"
    book.active.append(["ID", "TYPE", "File", "line", "check", "detail"])
    path = tmp_path / "rte.xlsx"
    book.save(path)

    screen, _ = peek.summarize(path)
    text = "\n".join(screen)

    assert "CodeMetrics 시트를 찾지 못했습니다" in text
    assert "RTE_Result" in text and "TYPE" in text


def test_width_counts_korean_as_two() -> None:
    assert peek.width("abc") == 3
    assert peek.width("가나") == 4
    assert peek.width(peek.fit("가나다라", 5)) <= 5
