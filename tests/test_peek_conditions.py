"""scripts/peek_conditions.py: 조건 시트 요약이 한 화면에 맞고 행·값 종류가 맞는지."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path
from types import ModuleType

from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "peek_conditions.py"


def _load() -> ModuleType:
    spec = importlib.util.spec_from_file_location("peek_conditions", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


peek = _load()

CATEGORIES = ["Mandatory", "Required", "Advisory"]


def _workbook(path: Path, misra_rules: int = 5) -> Path:
    book = Workbook()
    book.active.title = "Cover"
    book.active.append(["Polyspace 보고서"])
    for title in ("MISRA_Rule_Result", "CodeMetrics_Result", "RTE_Result"):
        book.create_sheet(title).append(["File", "Function", "line", "Check"])
    misra = book.create_sheet("MISRA_C_2012_Rule")
    misra.append(["MISRA C:2012 규칙"])  # 제목 행
    misra.append([])
    misra.append(["Rule", "Category", "Applied", "Description"])
    for i in range(misra_rules):
        misra.append([f"{i}.1", CATEGORIES[i % 3], "Yes" if i % 4 else "No", f"설명 {i} " * 20])
    metrics = book.create_sheet("Code Metrics")
    metrics.append(["Metric", "Threshold", "Direction"])
    metrics.append(["Cyclomatic Complexity", 10, "max"])
    metrics.append(["Comment Density", 20, "min"])
    book.create_sheet("RTE").append(["Check", "Target"])
    book["RTE"].append(["Overflow", "Red, Orange"])
    book.save(path)
    return path


def _block(screen: list[str], title: str) -> list[str]:
    """[title] 줄부터 다음 [시트] 줄 앞까지."""
    start = next(i for i, line in enumerate(screen) if line.startswith(f"[{title}]"))
    end = next((i for i in range(start + 1, len(screen)) if screen[i].startswith("[")), len(screen))
    return screen[start:end]


def test_finds_condition_sheets_only(tmp_path: Path) -> None:
    screen, _ = peek.summarize(_workbook(tmp_path / "p.xlsx"))
    text = "\n".join(screen)

    assert "조건 시트 3개" in screen[0]
    assert all(f"[{t}]" in text for t in ("MISRA_C_2012_Rule", "Code Metrics", "RTE"))
    assert "[MISRA_Rule_Result]" not in text and "[Cover]" not in text
    assert "Cover" in screen[1] and "RTE_Result" in screen[1]


def test_small_sheet_shows_every_row_with_row_numbers(tmp_path: Path) -> None:
    screen, _ = peek.summarize(_workbook(tmp_path / "p.xlsx"))
    block = _block(screen, "Code Metrics")

    assert "3행 중 앞 3행" in block[0] and "헤더 추정 1행" in block[0]
    assert any(
        line.split() == ["2:", "Cyclomatic", "Complexity", "|", "10", "|", "max"] for line in block
    )
    assert any(line.strip().startswith("3: Comment Density | 20 | min") for line in block)


def test_kinds_come_from_rows_under_the_guessed_header(tmp_path: Path) -> None:
    screen, _ = peek.summarize(_workbook(tmp_path / "p.xlsx", misra_rules=12))
    block = _block(screen, "MISRA_C_2012_Rule")

    assert "헤더 추정 3행" in block[0]
    assert any("Applied: Yes 9, No 3" in line for line in block)
    assert any("Category: Mandatory 4, Required 4, Advisory 4" in line for line in block)
    # 제목 1 + 헤더 1 + 규칙 12. 빈 행은 세지 않는다
    assert "14행 중" in block[0]


def test_big_sheet_fits_one_screen_and_file_has_everything(tmp_path: Path) -> None:
    screen, full = peek.summarize(_workbook(tmp_path / "p.xlsx", misra_rules=300))

    assert len(screen) <= peek.SCREEN_LINES - 1  # 마지막 한 줄은 파일 경로
    assert all(peek.width(line) <= peek.WIDTH for line in screen)
    # 작은 시트는 다 나오고 남은 줄은 큰 시트가 쓴다
    assert len(_block(screen, "Code Metrics")) == 1 + 3
    assert any("303: 299.1 | Advisory" in line for line in full)
    assert any(("설명 299 " * 20).strip() in line for line in full)


def test_many_condition_sheets_still_fit(tmp_path: Path) -> None:
    book = Workbook()
    book.active.title = "Cover"
    for n in range(12):
        sheet = book.create_sheet(f"MISRA part {n}")
        sheet.append(["Rule", "Category"])
        for i in range(20):
            sheet.append([f"{n}.{i}", CATEGORIES[i % 3]])
    path = tmp_path / "many.xlsx"
    book.save(path)

    screen, _ = peek.summarize(path)

    assert len(screen) <= peek.SCREEN_LINES - 1
    assert "조건 시트 12개" in screen[0]


def test_missing_sheets_explains_why(tmp_path: Path) -> None:
    book = Workbook()
    book.active.title = "Cover"
    book.active.append(["Polyspace", "보고서", "표지"])
    book.create_sheet("RTE_Result").append(["ID", "TYPE", "File"])
    path = tmp_path / "none.xlsx"
    book.save(path)

    screen, _ = peek.summarize(path)
    text = "\n".join(screen)

    assert "조건 시트를 찾지 못했습니다" in text
    assert "[Cover] Polyspace | 보고서 | 표지" in text and "[RTE_Result]" in text


def test_share_gives_small_needs_first() -> None:
    assert peek.share([3, 100, 50], 30) == [3, 14, 13]
    assert peek.share([2, 2], 30) == [2, 2]
    assert peek.share([5, 5, 5], 0) == [0, 0, 0]
