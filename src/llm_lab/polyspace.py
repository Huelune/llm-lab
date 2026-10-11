"""Polyspace 결과 엑셀에서 고른 종류(llm_lab.kinds)의 결과 시트를 읽는다."""

from __future__ import annotations

from collections.abc import Callable, Collection, Iterator
from dataclasses import dataclass
from typing import IO, TYPE_CHECKING, Any

from openpyxl import load_workbook

if TYPE_CHECKING:
    from llm_lab.kinds import Kind

# 헤더 위에 제목·요약 행이 있어도 찾도록 위에서부터 이만큼 훑는다
HEADER_SCAN_ROWS = 30


class SheetError(Exception):
    """고른 종류의 시트나 필수 열을 찾지 못함."""


@dataclass(frozen=True)
class Finding:
    row: int  # 엑셀 행 번호 (1부터)
    color: str  # RTE의 "Red" | "Orange". 다른 종류는 ""
    type: str  # TYPE 칸 원문
    file: str
    line: int | None  # 정수가 아니면 None
    check: str
    detail: str
    id: str = ""
    group: str = ""
    information: str = ""
    function: str = ""
    status: str = ""
    comment: str = ""
    # 종류별 값. 기본값이 있어야 DB에 저장된 예전 행 JSON도 그대로 읽힌다.
    category: str = ""  # MISRA의 실제 분류 (조건 시트 우선)
    note: str = ""  # 결과 카드에 보일 짧은 메모
    threshold: str = ""  # CodeMetrics: Threshold 칸 원문
    actual: str = ""  # CodeMetrics: Actual Value 칸 원문
    stage: str = ""  # CodeMetrics: 조건 시트의 판정 글


@dataclass(frozen=True)
class Prepared:
    """종류가 엑셀 전체에서 읽은 작업 정보 (조건 시트 등).

    data는 JSON으로 바꿀 수 있어야 한다. 작업에 저장해 두고 행 읽기와 LLM 메시지에 넘긴다.
    """

    data: Any = None
    note: str = ""  # 작업 메모에 남길 한 줄


@dataclass(frozen=True)
class Sheet:
    sheet: str
    findings: list[Finding]
    skipped: dict[str, int]  # 뺀 이유 → 행 수 (처음 나온 순서)
    prepared: Prepared = Prepared()


def norm(value: Any) -> str:
    """비교용: 공백을 한 칸으로 줄이고 대소문자를 무시한다."""
    return " ".join(str(value).split()).casefold() if value is not None else ""


def column_name(value: Any) -> str:
    """헤더 비교용: norm에 더해 끝의 콜론을 뗀다 ("Review Scope:")."""
    return norm(value).rstrip(":").strip()


def cell_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def line_number(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = cell_text(value)
    return int(text) if text.isdigit() else None


def find_header(sheet: Any, required: Collection[str]) -> tuple[int, dict[str, int]] | str:
    """(헤더 행 번호, 열 위치)를, 못 찾으면 오류 안내용으로 처음 보이는 행 내용을 돌려준다."""
    first_seen = ""
    rows = sheet.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True)
    for number, cells in enumerate(rows, start=1):
        columns: dict[str, int] = {}
        for index, cell in enumerate(cells):
            columns.setdefault(column_name(cell), index)
        if all(name in columns for name in required):
            return number, columns
        if not first_seen and any(cell is not None for cell in cells):
            first_seen = ", ".join(cell_text(cell) for cell in cells if cell is not None)
    return first_seen or "(비어 있음)"


def data_rows(
    sheet: Any, header_row: int, columns: dict[str, int]
) -> Iterator[tuple[int, Callable[[str], Any]]]:
    """헤더 아래의 (엑셀 행 번호, 열 이름 → 칸 값). 모든 칸이 빈 행은 건너뛴다."""
    rows = sheet.iter_rows(min_row=header_row + 1, values_only=True)
    for number, cells in enumerate(rows, start=header_row + 1):
        if all(cell_text(cell) == "" for cell in cells):
            continue

        def cell(name: str, cells: tuple[Any, ...] = cells) -> Any:
            index = columns.get(name)
            return cells[index] if index is not None and index < len(cells) else None

        yield number, cell


def read_sheet(source: IO[bytes], kind: Kind, selected: Collection[str]) -> Sheet:
    """엑셀에서 kind의 결과 시트를 찾아 분석할 행을 읽는다. selected는 고른 대상(배지 글자)이다."""
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise SheetError(f"엑셀(.xlsx) 파일로 읽을 수 없습니다: {exc}") from exc
    try:
        notes: list[str] = []
        for sheet in workbook.worksheets:
            title = sheet.title.casefold()
            if not title.endswith("_result"):
                notes.append(f"{sheet.title} (이름이 _Result로 끝나지 않음)")
            elif kind.sheet_word not in title:
                notes.append(f"{sheet.title} (이름에 {kind.sheet_word}가 없음)")
            elif isinstance(found := find_header(sheet, kind.required), str):
                notes.append(f"{sheet.title} (찾은 헤더: {found})")
            else:
                return _read_rows(sheet, *found, kind, selected, kind.prepare(workbook))
        raise SheetError(
            f"{kind.title} 시트를 찾지 못했습니다. 이름이 _Result로 끝나고 {kind.sheet_word}가 "
            f"들어가며 {', '.join(kind.required)} 열이 있는 시트가 필요합니다. "
            f"시트: {'; '.join(notes) or '없음'}"
        )
    finally:
        workbook.close()


def _read_rows(
    sheet: Any,
    header_row: int,
    columns: dict[str, int],
    kind: Kind,
    selected: Collection[str],
    prepared: Prepared,
) -> Sheet:
    findings: list[Finding] = []
    skipped: dict[str, int] = {}
    for number, cell in data_rows(sheet, header_row, columns):
        found = kind.read_row(cell, number, prepared.data)
        if isinstance(found, Finding):
            label = kind.label(found)
            if label in kind.choices and label not in selected:
                found = f"{label} (고르지 않음)"
        if isinstance(found, str):
            skipped[found] = skipped.get(found, 0) + 1
        else:
            findings.append(found)
    return Sheet(sheet.title, findings, skipped, prepared)
