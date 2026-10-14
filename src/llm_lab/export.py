"""작업 결과를 엑셀로 내보낸다 (행마다 한 줄: 엑셀 위치, 판단, 이유, diff).

회사 PC 안에서 보고 나누는 파일이다. 화면 캡처용 요약과 달리 이름·이유·코드가 그대로 들어간다.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from llm_lab.kinds import KINDS
from llm_lab.pages import KIND_LABELS, _edits_other_files, _explain, _kind
from llm_lab.store import Job, Row

MAX_CELL = 32767  # 엑셀 칸 하나에 들어가는 글자 수
COLUMNS = (  # (머리글, 폭)
    ("엑셀 행", 8),
    ("대상", 12),
    ("검사", 30),
    ("파일", 28),
    ("줄", 6),
    ("함수", 22),
    ("판단", 12),
    ("다른 파일 수정", 8),
    ("이유 / 할 일", 60),
    ("diff", 80),
)


def _text(value: object) -> object:
    """칸에 넣을 값. 엑셀이 못 받는 제어 문자는 빼고, 너무 긴 글은 자른다."""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value if len(value) <= MAX_CELL else value[: MAX_CELL - 20] + "\n… (잘림)"


def _reason(row: Row) -> str:
    if row.result is not None:
        return row.result.reason
    if _kind(row) == "problem":
        summary, action, _ = _explain(row)
        return summary + (f"\n할 일: {action}" if action else "")
    return ""


def export_workbook(job: Job, rows: list[Row]) -> bytes:
    kind = KINDS[job.kind]
    book = Workbook()
    sheet = book.active
    sheet.title = "결과"
    sheet.append([title for title, _ in COLUMNS])
    for row in rows:
        finding = row.finding
        fix = row.result is not None and row.result.decision == "fix"
        values = (
            finding.row,
            kind.label(finding),
            finding.check,
            finding.file,
            finding.line,
            finding.function,
            KIND_LABELS[_kind(row)],
            "예" if fix and _edits_other_files(row) else "",
            _reason(row),
            row.result.diff if row.result is not None else "",
        )
        sheet.append([_text(value) for value in values])
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            # =, +, - 로 시작하는 이유·diff가 수식으로 읽히지 않게 글자로 둔다
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF0F3")
    for index, (_, width) in enumerate(COLUMNS):
        sheet.column_dimensions[sheet.cell(1, index + 1).column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
